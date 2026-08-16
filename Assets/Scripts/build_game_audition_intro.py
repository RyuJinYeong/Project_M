from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE


OUTPUT = Path(r"D:\Project\Project_M\Assets\Scripts\2026_경기게임오디션_게임소개서_제출본.docx")
IMG_NIGHT = Path(r"C:\Users\G9\AppData\Local\Temp\codex-clipboard-96bc6eef-e749-4141-b6a4-d527415cb353.png")
IMG_STATUS = Path(r"C:\Users\G9\AppData\Local\Temp\codex-clipboard-b70b2420-7359-4c4a-b436-91c76e2c9035.png")
IMG_VOTE = Path(r"C:\Users\G9\AppData\Local\Temp\codex-clipboard-35b26e14-b401-4b67-8b29-c0bcb42cda90.png")

FONT = "Malgun Gothic"
NAVY = "12263A"
TEAL = "2A9D8F"
GOLD = "D6A84B"
RED = "A44747"
INK = "1B2A34"
MUTED = "65737E"
PALE = "EEF3F5"
PALE_GOLD = "F7F0DF"
WHITE = "FFFFFF"
BLACK = "000000"
WIDTH_DXA = 9360
TABLE_INDENT = 120


def set_run(run, size=None, bold=None, color=INK, italic=None, font=FONT):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(element, fill):
    pr = element.get_or_add_tcPr() if hasattr(element, "get_or_add_tcPr") else element.get_or_add_pPr()
    shd = pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="CFD8DE", size=6, inside=True):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    names = ["top", "left", "bottom", "right"] + (["insideH", "insideV"] if inside else [])
    for name in names:
        edge = borders.find(qn(f"w:{name}"))
        if edge is None:
            edge = OxmlElement(f"w:{name}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), str(size))
        edge.set(qn("w:color"), color)


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text, size=9.2, bold=False, color=INK, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    set_run(r, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths, header_fill=NAVY, font_size=9.0):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for i, h in enumerate(headers):
        shade(table.rows[0].cells[i]._tc, header_fill)
        set_cell_text(table.rows[0].cells[i], h, size=9.2, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
    for ridx, row_data in enumerate(rows):
        cells = table.add_row().cells
        for i, val in enumerate(row_data):
            if ridx % 2:
                shade(cells[i]._tc, "F7F9FA")
            set_cell_text(cells[i], str(val), size=font_size, bold=(i == 0), color=INK)
    return table


def add_bullet_numbering(doc):
    numbering = doc.part.numbering_part.element
    existing_abs = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    existing_num = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abs_id = max(existing_abs, default=0) + 1
    num_id = max(existing_num, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abs_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start"); start.set(qn("w:val"), "1"); lvl.append(start)
    num_fmt = OxmlElement("w:numFmt"); num_fmt.set(qn("w:val"), "bullet"); lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText"); lvl_text.set(qn("w:val"), "•"); lvl.append(lvl_text)
    jc = OxmlElement("w:lvlJc"); jc.set(qn("w:val"), "left"); lvl.append(jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs"); tab = OxmlElement("w:tab"); tab.set(qn("w:val"), "num"); tab.set(qn("w:pos"), "540"); tabs.append(tab); p_pr.append(tabs)
    ind = OxmlElement("w:ind"); ind.set(qn("w:left"), "540"); ind.set(qn("w:hanging"), "280"); p_pr.append(ind)
    spacing = OxmlElement("w:spacing"); spacing.set(qn("w:after"), "80"); spacing.set(qn("w:line"), "290"); spacing.set(qn("w:lineRule"), "auto"); p_pr.append(spacing)
    lvl.append(p_pr)
    r_pr = OxmlElement("w:rPr"); fonts = OxmlElement("w:rFonts"); fonts.set(qn("w:ascii"), FONT); fonts.set(qn("w:hAnsi"), FONT); fonts.set(qn("w:eastAsia"), FONT); r_pr.append(fonts); lvl.append(r_pr)
    abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num"); num.set(qn("w:numId"), str(num_id)); abs_node = OxmlElement("w:abstractNumId"); abs_node.set(qn("w:val"), str(abs_id)); num.append(abs_node); numbering.append(num)
    return num_id


def bullet(doc, text, num_id, color=INK, bold_prefix=None, size=10.2):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.208
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId"); num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num]); p_pr.append(num_pr)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix); set_run(r1, size=size, bold=True, color=color)
        r2 = p.add_run(text[len(bold_prefix):]); set_run(r2, size=size, color=color)
    else:
        r = p.add_run(text); set_run(r, size=size, color=color)
    return p


def add_para(doc, text, size=10.3, bold=False, color=INK, align=WD_ALIGN_PARAGRAPH.LEFT, after=6, italic=False, keep=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.keep_together = keep
    r = p.add_run(text)
    set_run(r, size=size, bold=bold, color=color, italic=italic)
    return p


def add_heading(doc, text, level=1, kicker=None):
    if kicker:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(kicker.upper())
        set_run(r, size=8.5, bold=True, color=TEAL)
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    return p


def add_page_title(doc, number, title, subtitle):
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = True
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f"{number:02d}")
    set_run(r, size=9, bold=True, color=TEAL)
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(5)
    p.add_run(title)
    add_para(doc, subtitle, size=10.5, color=MUTED, after=12)


def add_callout(doc, label, text, fill=PALE_GOLD, accent=GOLD):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [WIDTH_DXA])
    set_table_borders(table, color=accent, size=8, inside=False)
    cell = table.cell(0, 0)
    shade(cell._tc, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label + "  ")
    set_run(r, size=9.2, bold=True, color=accent)
    r = p.add_run(text)
    set_run(r, size=10.2, bold=True, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_screenshot(doc, path, caption, top_crop=13000, bottom_crop=7000, width=6.5, height=4.0):
    if not path.exists():
        add_callout(doc, "이미지 자리", caption, fill="F7F9FA", accent=MUTED)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run()
    shape = run.add_picture(str(path), width=Inches(width), height=Inches(height))
    shape._inline.docPr.set("descr", caption)
    shape._inline.docPr.set("title", caption)
    blip_fill = shape._inline.graphic.graphicData.pic.blipFill
    src_rect = OxmlElement("a:srcRect")
    src_rect.set("t", str(top_crop))
    src_rect.set("b", str(bottom_crop))
    blip_fill.insert(1, src_rect)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    r = cap.add_run(caption)
    set_run(r, size=8.5, color=MUTED, italic=True)


def page_break(doc):
    # Section openings use page_break_before on their title paragraph.
    # This avoids blank pages when the preceding section exactly fills a page.
    return


def setup_styles(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.25
    for name, size, color, before, after in [
        ("Heading 1", 16, NAVY, 16, 8),
        ("Heading 2", 12.5, TEAL, 12, 6),
        ("Heading 3", 11.2, NAVY, 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    r = hp.add_run("[[GAME_TITLE]]  |  2026 경기게임오디션 일반부문")
    set_run(r, size=8.2, color=MUTED)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    r = fp.add_run("[[TEAM_NAME]]   •   ")
    set_run(r, size=8, color=MUTED)
    fld_begin = OxmlElement("w:fldChar"); fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar"); fld_sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t"); txt.text = "1"
    fld_end = OxmlElement("w:fldChar"); fld_end.set(qn("w:fldCharType"), "end")
    r2 = fp.add_run(); r2._r.extend([fld_begin, instr, fld_sep, txt, fld_end]); set_run(r2, size=8, color=MUTED)


def build():
    doc = Document()
    setup_styles(doc)
    bullet_id = add_bullet_numbering(doc)

    # Cover — proposal_centerpiece header pattern.
    add_para(doc, "2026 경기게임오디션 · 일반부문", size=10.5, bold=True, color=TEAL, align=WD_ALIGN_PARAGRAPH.CENTER, after=34)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("[[GAME_TITLE]]")
    set_run(r, size=30, bold=True, color=NAVY)
    add_para(doc, "낮에는 말로 의심하고, 밤에는 행동으로 증명한다", size=15, bold=True, color=GOLD, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    add_para(doc, "최대 12인 3D 멀티플레이 소셜 디덕션", size=12, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=30)
    add_callout(doc, "한 줄 소개", "정체가 지워진 영체들이 밤의 마을에 남긴 도구·동선·침입의 흔적을, 다음 날 토론과 투표로 해석하는 공간 기반 소셜 디덕션 게임.")
    add_para(doc, "", after=20)
    meta = doc.add_table(rows=4, cols=2)
    set_table_geometry(meta, [4680, 4680])
    set_table_borders(meta, color="D8E0E5", size=5)
    pairs = [
        ("개발팀", "[[TEAM_NAME]]"),
        ("개발 형태", "1인 개발"),
        ("플랫폼", "PC (Windows) · Steam 출시 목표"),
        ("개발 단계", "핵심 루프가 연결된 플레이어블 멀티플레이 프로토타입"),
    ]
    for i, (a, b) in enumerate(pairs):
        set_cell_text(meta.cell(i, 0), a, size=9.2, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(meta.cell(i, 1), b, size=9.2, color=INK, align=WD_ALIGN_PARAGRAPH.CENTER)
        if i % 2 == 0:
            shade(meta.cell(i, 0)._tc, PALE); shade(meta.cell(i, 1)._tc, PALE)
    add_para(doc, "게임명과 팀명은 문서 전체에서 각각 [[GAME_TITLE]], [[TEAM_NAME]]을 찾기·바꾸기로 일괄 교체할 수 있습니다.", size=8.2, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)

    page_break(doc)
    add_page_title(doc, 1, "프로젝트 요약", "익숙한 낮의 토론에, 직접 움직이는 밤의 행동과 공간 증거를 결합합니다.")
    add_callout(doc, "CORE QUESTION", "“어젯밤 내가 본 그 영체는 누구였고, 왜 그 집에 들어갔는가?”")
    add_para(doc, "[[GAME_TITLE]]은 플레이어가 저주받은 마을의 주민, 마녀, 또는 독자적인 목표를 가진 이방인이 되어 서로의 정체를 추리하는 게임입니다. 낮에는 공개 토론과 추방 투표가 진행되고, 밤에는 육체에서 빠져나온 영체로 마을을 돌아다니며 직업 행동·직무·침입·살해·방어를 직접 수행합니다.", size=10.4)
    add_para(doc, "밤이 되면 닉네임과 외형이 사라지고, 손에 든 도구와 이동한 방향, 출입한 문, 행동 모션과 소리만 남습니다. 마녀와 일부 이방인은 주민 도구로 위장할 수 있어 단일 단서가 정답이 되지 않습니다. 플레이어는 자신이 본 행동을 아침의 진술과 결과에 맞춰보며 신뢰와 거짓을 가려냅니다.", size=10.4)
    add_heading(doc, "기본 정보", level=2)
    add_table(doc, ["항목", "내용"], [
        ("장르", "3D 멀티플레이 소셜 디덕션"),
        ("플레이 인원", "최대 12인"),
        ("엔진·네트워크", "Unity 6 · Netcode for GameObjects · Unity Relay · Vivox"),
        ("핵심 콘텐츠", "3개 진영 · 20종 역할 · 집 기반 침입과 방어 · 야간 직무와 마을 안정도"),
        ("현재 단계", "로비부터 게임 결과까지 실제 다인 플레이가 가능한 프로토타입"),
    ], [2300, 7060], font_size=9.2)
    add_heading(doc, "세 가지 차별점", level=2)
    add_table(doc, ["차별점", "게임플레이 효과"], [
        ("밤의 행동이 증거", "선택창 대신 마을을 직접 이동하고, 그 과정이 목격 가능한 단서로 남음"),
        ("집이 개인 던전", "침입·방어·감금·살해가 각자의 집과 정문·뒷문에서 발생"),
        ("모든 진영의 야간 동선", "주민 직무·마녀 교란·이방인 목표가 매일 다른 알리바이를 생성"),
    ], [2850, 6510], font_size=8.9)

    page_break(doc)
    add_page_title(doc, 2, "핵심 게임 루프", "아침의 말과 밤의 행동이 서로의 원인과 결과가 되는 반복 구조입니다.")
    phase_rows = [
        ("아침 토론·투표", "목격 정보 공유, 공개 음성·텍스트 토론, 추방 또는 기권"),
        ("밤 준비", "마녀 살해 담당 선정, 일부 역할의 사전 선택"),
        ("밤 행동", "영체 이동, 직업 행동, 직무, 교란, 침입과 방어"),
        ("밤 결과", "사망·보호·역할 전환·직무 수행률·안정도 반영"),
    ]
    add_table(doc, ["페이즈", "플레이어가 하는 일"], phase_rows, [2650, 6710], font_size=9.2)
    add_heading(doc, "아침 — 말의 신뢰도를 검증하는 시간", level=2)
    add_para(doc, "생존자는 지난밤 자신이 직접 본 도구·동선·출입과 직업 능력의 결과를 근거로 토론합니다. 대화와 투표는 하나의 60초 시간 안에서 동시에 진행되고, 모든 생존자가 투표를 확정하면 남은 시간은 5초로 줄어듭니다. 투표 결과는 현황판 기록으로 남아 다음 라운드의 정치적 증거가 됩니다.")
    add_heading(doc, "밤 — 말의 근거를 직접 만드는 시간", level=2)
    add_para(doc, "육체가 집에 남아 있는 동안 플레이어는 영체로 움직입니다. 주민은 직업 행동과 매일 두 개의 직무를, 마녀는 살해 담당과 교란 공작을 나누어 수행합니다. 능력을 쓰려면 대상의 집이나 문 앞에 직접 접근해야 하므로 성공 가능성과 목격·전투 위험을 함께 계산해야 합니다.")

    page_break(doc)
    add_page_title(doc, 3, "밤의 익명성과 공간 기반 추리", "정체 정보는 사라지고, 행동의 흔적만 남습니다.")
    add_callout(doc, "DESIGN PRINCIPLE", "도구 하나로 역할을 확정할 수 없게 하고, 도구 + 위치 + 시간 + 결과가 함께 맞아야 믿을 수 있는 증거가 되게 합니다.")
    add_table(doc, ["밤에 숨겨지는 정보", "밤에 관찰할 수 있는 정보"], [
        ("닉네임·얼굴·의상·캐릭터 색", "손에 든 역할 도구와 위장 도구"),
        ("머리 위 신원 표시", "정문·뒷문 출입 위치와 이동 방향"),
        ("낮의 공개 대화", "직업 행동·직무·교란 공작의 모션과 소리"),
        ("다른 플레이어의 정확한 역할", "제압·감금·시체·저주 인형·집 상태 변화"),
    ], [4300, 5060], font_size=9.2)
    add_heading(doc, "정보량도 매일 달라집니다", level=2)
    bullet(doc, "횃불을 켠 생존자는 야간 위치 음성에 참여하고, 과반수가 점등하면 마을 전체가 밝아집니다.", bullet_id)
    bullet(doc, "마을 안정도가 낮아지면 주민과 이방인의 원거리 영체 식별 거리와 타인의 집 이름표가 단계적으로 제한됩니다.", bullet_id)
    bullet(doc, "마녀·자유 관전자·최종 결투 참가자는 안정도 시야 제한에서 제외되어, 정보 비대칭이 진영 전략으로 이어집니다.", bullet_id)

    page_break(doc)
    add_page_title(doc, 4, "집 침입·실시간 저항·사망 이후", "전투는 처치 경쟁이 아니라 상대의 행동권과 알리바이를 흔드는 수단입니다.")
    add_heading(doc, "집은 밤의 핵심 플레이 공간", level=2)
    add_para(doc, "각 플레이어의 육체는 밤 동안 자신의 집에 남습니다. 조사·보호·살해 능력은 육체나 집의 정문 같은 실제 대상과 연결되고, 침입자는 목적지까지 이동해 행동 게이지를 채워야 합니다. 피격되면 진행 중인 직업 행동과 직무가 즉시 취소되며, 피격 직후에는 짧은 출입 제한이 걸려 문을 이용한 즉시 도주를 막습니다.")
    add_table(doc, ["상황", "3회 피격 시 결과", "플레이 의미"], [
        ("자기 집", "잠시 제압", "침입자에게 짧은 행동 기회가 생김"),
        ("타인의 집", "그 밤 동안 감금", "무리한 침입이 큰 행동권 손실로 이어짐"),
        ("집 밖", "밀쳐내기·이동 방해", "추적과 견제 중심의 야외 전투 유지"),
        ("주정뱅이", "야외에서도 제압", "집 밖 취침이라는 역할 개성을 전투 규칙에 반영"),
    ], [1900, 2400, 5060], font_size=8.9)
    add_heading(doc, "사망자는 곧바로 모든 정보를 얻지 못합니다", level=2)
    add_para(doc, "사망 직후에는 영매사의 교신 대상이 될 수 있는 제한 관전 상태가 유지되어, 다른 플레이어의 역할과 전체 동선을 볼 수 없습니다. 해당 기간이 끝난 뒤에야 자유 관전자로 전환됩니다. 추방자는 시체 없이 도구만 남도록 구현되어 있으며, 퇴마 대상의 검시 결과와 영매사 교신 차단은 최종 회귀 테스트 항목으로 관리합니다.")
    add_heading(doc, "대화 채널도 게임 상태에 맞춰 분리됩니다", level=2)
    add_table(doc, ["채널", "참여자", "용도"], [
        ("아침 공개 채팅", "모든 생존자", "토론과 투표 설득"),
        ("야간 마녀 채팅", "마녀 진영", "살해 담당·교란 동선 협의"),
        ("영매사 교신", "영매사와 지정 사망자", "제한된 사후 정보 교환"),
        ("야간 위치 음성", "횃불 조건을 충족한 생존자", "가까운 플레이어와 위험한 현장 소통"),
    ], [2350, 2650, 4360], font_size=8.9)

    page_break(doc)
    add_page_title(doc, 5, "세 진영과 20종 역할", "역할 수보다 역할 간 충돌과 공간 상호작용을 우선합니다.")
    add_table(doc, ["진영", "표시명", "목표", "플레이 성격"], [
        ("Citizen", "주민", "마녀와 생존을 위협하는 이방인 제거", "관찰·보호·조사·직무·토론"),
        ("Mafia", "마녀", "마녀를 제외한 모든 생존자 제거", "위장·살해 담당·교란·침입"),
        ("Neutral", "이방인", "역할별 독립 승리 조건 달성", "계승·독립 살해·추방 유도"),
    ], [1500, 1450, 3300, 3110], font_size=8.8)
    add_heading(doc, "주민 13종", level=2)
    add_para(doc, "농부 · 잡역부 · 대장장이 · 집행관 · 의원 · 수사관 · 주정뱅이 · 퇴마사 · 검시관 · 영매사 · 사냥꾼 · 장의사 · 행상인", size=10.3, bold=True, color=TEAL)
    add_para(doc, "직업은 보호·조사·추적·사후 정보·공간 통제·도구 변화로 역할이 나뉩니다. 강한 정보를 버튼 한 번으로 주기보다, 대상에게 접근하고 행동 시간을 확보해야 결과를 얻도록 설계했습니다.")
    add_heading(doc, "마녀 4종", level=2)
    add_table(doc, ["역할", "비살해 플레이", "살해 담당일 때"], [
        ("마녀", "주민 도구 위장·교란 공작", "육체 또는 정문에 독 사용"),
        ("저주술사", "주민 도구 위장·저주 인형 운반", "대상의 설치점에 저주 인형 배치"),
        ("내통자", "부여된 주민 역할 능력 사용", "직접 살해와 주민 능력 중 선택"),
        ("잠입자", "뒷문·뒷골목 침입", "육체에 직접 살해"),
    ], [1900, 3900, 3560], font_size=8.8)
    add_heading(doc, "이방인 3종", level=2)
    add_table(doc, ["역할", "핵심 규칙", "승리 방향"], [
        ("도둑", "사망자가 떨어뜨린 도구를 훔쳐 역할과 진영 계승", "미전환 도둑만 남으면 공동 승리"),
        ("살인귀", "매 밤 직접 살해·3회 피격 전투", "마지막 생존 또는 최종 결투 승리"),
        ("순교자", "주민 도구로 위장하며 추방을 유도", "아침 투표로 추방되면 개인 승리"),
    ], [1800, 4340, 3220], font_size=8.8)

    page_break(doc)
    add_page_title(doc, 6, "마녀 팀플레이와 저주 인형", "한 번의 팀 살해권을 누가, 어떤 위험을 감수하고 사용할지 매일 결정합니다.")
    add_heading(doc, "살해 담당 선택", level=2)
    add_para(doc, "마녀 진영은 밤 준비 단계에서 그날의 살해 담당 한 명을 정합니다. 담당 역할에 따라 독살·저주 인형·직접 살해로 방법이 달라지고, 팀 살해는 한 번만 소비됩니다. 살해 담당이 아닌 마녀는 자신의 비살해 능력이나 교란 공작으로 알리바이와 침입 기회를 만듭니다.")
    add_table(doc, ["담당 역할", "살해 방식", "강점과 위험"], [
        ("마녀", "독살", "의원의 보호에 막힐 수 있어 접근 위치와 타이밍이 중요"),
        ("저주술사", "저주 인형", "보호를 무시하지만 누구나 옮길 수 있어 아군에게 되돌아올 수 있음"),
        ("내통자", "직접 살해", "그날의 주민 능력 사용을 포기해야 함"),
        ("잠입자", "직접 살해", "뒷문 접근이 유리하지만 그 동선 자체가 단서"),
    ], [1900, 1900, 5560], font_size=8.9)
    add_heading(doc, "저주 인형 — 살해 수단이면서 모두가 만질 수 있는 사건", level=2)
    bullet(doc, "최초 설치에는 다른 살해 행동과 같은 3초 행동 시간이 필요해, 목격과 피격으로 취소될 수 있습니다.", bullet_id)
    bullet(doc, "설치된 인형은 E키로 집어 다른 유효 설치점으로 옮길 수 있고, 다른 플레이어도 운반할 수 있습니다.", bullet_id)
    bullet(doc, "날이 밝을 때 인형이 가리키는 대상이 사망하며 의원 보호를 무시하고 진영도 가리지 않습니다.", bullet_id)
    bullet(doc, "운반자가 감금·제압되거나 야외에서 3회 피격되면 인형은 유효한 자신의 설치점으로 돌아갑니다.", bullet_id)
    bullet(doc, "첫 설치 전에 저주술사가 감금되면 인형은 자동 설치되지 않아, 조기 노출과 구조를 악용하는 예외를 차단합니다.", bullet_id)
    add_callout(doc, "상황 생성", "주민은 인형을 옮겨 살해를 막거나 위험을 되돌릴 수 있고, 마녀는 인형의 이동을 추적하며 팀의 살해 계획을 다시 세워야 합니다.", fill=PALE, accent=TEAL)

    page_break(doc)
    add_page_title(doc, 7, "야간 직무·마을 안정도·교란 공작", "반복 시스템이 모든 진영의 이동 이유와 장기적인 정보 환경을 만듭니다.")
    add_heading(doc, "주민의 야간 직무", level=2)
    add_para(doc, "생존 주민은 매일 밤 두 개의 직무를 무작위로 배정받습니다. 지점은 광장 우물과 각 집의 정문·뒷문 잠금장치, 침대 프레임, 외부 벤치입니다. 직업 행동과 대상이 겹치면 R키로 ‘직무 우선 / 직업 능력 우선’을 바꿀 수 있습니다.")
    add_table(doc, ["구역", "직무 지점", "추리에 주는 효과"], [
        ("마을 중앙", "광장 우물", "여러 집의 동선이 겹치는 공개 목격 지점"),
        ("집 외부", "정문·뒷문·외부 벤치", "침입·잠복·직무 동선이 겹치는 오해 지점"),
        ("집 내부", "침대 프레임", "집주인 방어와 타인 직무가 충돌하는 위험 지점"),
    ], [1900, 3000, 4460], font_size=8.9)
    add_heading(doc, "안정도는 전날 결과를 누적합니다", level=2)
    bullet(doc, "초기 안정도는 75%이며, 직무 완벽 수행 시 최대 +25%p, 모두 실패 시 최대 -25%p가 반영됩니다.", bullet_id)
    bullet(doc, "미완료 직무 1개당 다음 주민 능력 실패 확률이 25%p 증가하고, 직무 패널티는 최대 50%p입니다.", bullet_id)
    bullet(doc, "불안정 패널티는 주민 능력에 최대 20%p까지 추가되고, 주민·이방인의 원거리 영체와 이름표 식별도 줄어듭니다.", bullet_id)
    bullet(doc, "마녀는 안정도 능력 패널티를 받지 않지만 안정도 정보는 동일하게 확인해 주민처럼 행동할 수 있습니다.", bullet_id)
    add_heading(doc, "마녀 교란 공작", level=2)
    add_para(doc, "살해 담당이 아닌 마녀는 두 후보 중 하나의 교란 공작을 골라 수행합니다. 혼자 남은 마녀는 살해 담당이면서도 교란을 수행할 수 있습니다. 성공한 교란은 주민의 유효 직무 완료 수를 상쇄하고, 주민에게는 교란 횟수를 직접 공개하지 않은 채 직무 수행률과 안정도 변화로만 나타납니다.")

    page_break(doc)
    add_page_title(doc, 8, "이방인과 특수 종료", "같은 밤을 공유하지만, 승리의 방향은 주민과 마녀 어느 쪽에도 고정되지 않습니다.")
    add_heading(doc, "도둑 — 사망자의 도구를 훔쳐 운명을 바꿉니다", level=2)
    add_para(doc, "도둑은 쇠지렛대를 들고 시작하고, 바닥에 떨어진 역할 도구를 획득하면 해당 역할과 진영을 그대로 계승합니다. 마녀나 다른 이방인의 도구도 예외가 없습니다. 전환하지 않은 도둑들만 남으면 공동 승리하며, 도구를 훔치지 않은 도둑이 주민·순교자와 함께 남은 상황은 별도 승리 판정으로 성급히 종료되지 않습니다.")
    add_heading(doc, "살인귀 — 야간 직접 살해와 최종 결투", level=2)
    add_para(doc, "살인귀는 잠입자의 뒷문 특화가 없는 대신 매 밤 육체에 F키 행동을 사용해 직접 살해합니다. 주민과 1대1이면 살인귀 승리, 마녀와 1대1이면 일반 페이즈를 중단하고 광장 결투로 전환됩니다.")
    add_table(doc, ["최종 결투 규칙", "적용 내용"], [
        ("시작 연출", "중앙 안내 후 각자의 집 정문 밖에서 영체로 시작"),
        ("허용 행동", "직업 능력과 문 상호작용 잠금, 기존 위장 도구 기본 공격만 허용"),
        ("승패", "3회 피격된 참가자를 즉시 최종 사망 처리하고 생존 진영 승리"),
    ], [2700, 6660], font_size=9.0)
    add_heading(doc, "순교자 — 추방을 목표로 하는 심리전", level=2)
    add_para(doc, "순교자는 농부·잡역부·대장장이 도구 중 하나로 위장합니다. 밤에 살해되거나 다른 방식으로 제거되는 것은 승리가 아니며, 오직 아침 투표에서 추방되었을 때 정체와 개인 승리를 공개하고 경기를 종료합니다. 따라서 순교자가 남아 있는 종반에는 단순 인원수만으로 승패를 확정할 수 없습니다.")
    add_callout(doc, "반복 플레이", "역할 조합, 살해 담당, 위장 도구, 두 개의 직무, 교란 공작, 안정도와 횃불 상태가 동시에 달라져 같은 역할도 매번 다른 동선을 선택하게 됩니다.", fill=PALE, accent=TEAL)

    page_break(doc)
    add_page_title(doc, 9, "현재 구현 상태와 기술 기반", "아이디어 설명 단계가 아니라, 로비부터 결과까지 이어지는 플레이어블 프로토타입입니다.")
    add_table(doc, ["영역", "현재 빌드 상태"], [
        ("멀티플레이 기반", "Unity NGO 호스트·클라이언트, Relay 로비, 재접속·로비 복귀 흐름, 최대 12인 설정"),
        ("매치 루프", "아침 투표, 밤 준비·행동·결과, 특수 승리·최종 결투, 결과 역할 공개"),
        ("공간·전투", "집 배정, 정문·뒷문, 영체 이동, 3회 피격 제압·감금, 출입 제한, 야간 HP 유지"),
        ("역할", "주민 13종, 마녀 4종, 이방인 3종의 배정·행동·승리 규칙"),
        ("장기 시스템", "직무 2개, 안정도 0~100, 능력 실패·식별 거리·횃불 보상, 마녀 교란"),
        ("커뮤니케이션", "Vivox 음성, 아침 공개 채팅, 마녀 야간 채팅, 영매사 교신, 다중 줄 텍스트 UI"),
        ("UX·기록", "도움말, 상태이상 HUD, 개인 HP, TAB 현황판, 공개 기록·개인 행동 기록"),
    ], [2300, 7060], font_size=8.75)
    add_heading(doc, "서버 권한 구조", level=2)
    add_para(doc, "역할 배정, 살해·보호·감금·사망·승리 판정은 서버 권한으로 처리하고, 클라이언트는 입력과 표시를 담당합니다. 고지연 접속자가 직업 배정을 받기 전에 게임으로 넘어가지 않도록 참가자 준비 완료를 확인한 뒤 매치를 시작하며, 이동·타격에는 제한적인 보간과 서버 검증을 적용합니다.")
    add_heading(doc, "남은 위험과 대응", level=2)
    add_table(doc, ["위험", "대응 계획"], [
        ("6~12인 실제 인원 테스트 부족", "외부 플레이테스트와 오디션 시연으로 역할·정보량·라운드 길이 검증"),
        ("고지연 Relay·Vivox 환경", "150~200ms 환경 회귀 테스트, 준비 완료 게이트와 재접속 흐름 반복 검증"),
        ("20종 역할의 예외 조합", "테스트 목록 기반으로 승리·사망·도구 계승·교신 우선순위 검증"),
        ("1인 개발 범위", "새 시스템 확장보다 밸런스·온보딩·안정성·연출 폴리싱 우선"),
    ], [3300, 6060], font_size=8.8)

    page_break(doc)
    add_page_title(doc, 10, "현재 개발 빌드 화면", "기획의 핵심 정보 구조가 실제 멀티플레이 UI와 게임플레이에 연결되어 있습니다.")
    add_screenshot(doc, IMG_VOTE, "아침 추방 투표 — 공개 토론과 투표, 사망·진영·관전 상태별 정보 표시", top_crop=12000, bottom_crop=9000, height=2.1)
    add_screenshot(doc, IMG_NIGHT, "야간 영체 시점 — 진영 전용 채팅, 개인 HP, 직무와 안정도 UI", top_crop=13000, bottom_crop=9000, height=2.1)
    add_screenshot(doc, IMG_STATUS, "TAB 마을 현황판 — 생존 현황, 직무 수행률, 안정도, 횃불과 개인 행동 기록", top_crop=13000, bottom_crop=9000, height=2.1)

    page_break(doc)
    add_page_title(doc, 11, "재미 요소와 시장성", "친숙한 소셜 디덕션 문법 위에, 관찰 가능한 밤의 사건을 더합니다.")
    add_heading(doc, "목표 이용자", level=2)
    add_table(doc, ["타깃", "기대 가치"], [
        ("친구 그룹의 파티·추리 이용자", "낮의 대화와 밤의 직접 행동을 번갈아 즐기는 협력·배신 경험"),
        ("마피아·역할 추리 장르 이용자", "도구·동선·문·시간이 만드는 새로운 증거 층위"),
        ("스트리머·콘텐츠 크리에이터", "침입·오해·추격·저주 인형 역전처럼 설명 가능한 사건이 매 라운드 발생"),
    ], [3500, 5860], font_size=9.0)
    add_heading(doc, "장르 안에서의 포지셔닝", level=2)
    add_table(doc, ["구분", "전통 역할 추리", "행동형 파티 추리", "[[GAME_TITLE]]"], [
        ("낮의 토론·투표", "핵심", "핵심", "핵심"),
        ("밤의 직접 이동", "선택 UI 중심", "상시 이동", "익명 영체로 페이즈화"),
        ("개인 집 침입·방어", "제한적", "제한적", "핵심 공간 규칙"),
        ("장기 마을 상태", "제한적", "작업 중심", "직무·안정도·횃불로 정보 환경 변화"),
        ("도구 위장과 동선 증거", "역할 정보 중심", "행동 증거", "위장 가능 도구 + 위치·시간 조합"),
    ], [1900, 2100, 2100, 3260], font_size=8.15)
    add_heading(doc, "방송과 반복 플레이에 적합한 이유", level=2)
    bullet(doc, "누군가의 집에서 벌어진 몸싸움, 잘못 본 도구, 되돌아온 저주 인형처럼 ‘말로 설명 가능한 사건’이 자연스럽게 생깁니다.", bullet_id)
    bullet(doc, "말하기에 자신이 없는 플레이어도 관찰·추적·직무·방어 행동으로 팀에 기여할 수 있습니다.", bullet_id)
    bullet(doc, "역할 암기보다 그날 밤의 이동과 선택이 결과를 바꾸므로, 같은 조합에서도 새로운 이야기가 발생합니다.", bullet_id)
    page_break(doc)
    add_page_title(doc, 12, "개발 로드맵과 오디션 활용 계획", "새 기반을 더 만드는 단계보다, 실제 다인 환경에서 재미와 완성도를 끌어올리는 단계입니다.")
    add_table(doc, ["단계", "핵심 목표", "완료 기준"], [
        ("1. 다인 회귀 테스트", "20종 역할·승리·사망·도구 예외 검증", "6~12인 테스트에서 치명적 진행 중단 제거"),
        ("2. 네트워크·음성 안정화", "Relay, 150~200ms, Vivox 채널 검증", "고지연 참가자의 정상 배정·이동·교신"),
        ("3. 밸런싱·온보딩", "직무·안정도·역할 조합·도움말 조정", "첫 플레이어가 핵심 루프를 스스로 수행"),
        ("4. 데모 폴리싱", "사운드·연출·그래픽 옵션·최적화", "외부 시연용 안정 빌드 완성"),
        ("5. 출시 준비", "Steam 페이지·빌드 파이프라인·운영", "배포와 피드백 반영이 가능한 구조"),
    ], [2100, 3760, 3500], font_size=8.55)
    add_heading(doc, "경기게임오디션에서 검증할 네 가지", level=2)
    add_table(doc, ["심사 관점", "검증 질문"], [
        ("기획의 우수성", "야간 익명 이동·집 중심 추리·직무와 안정도가 하나의 루프로 자연스럽게 이어지는가?"),
        ("완성도", "6~12인 환경에서 NGO·Relay·Vivox와 20종 역할 상호작용이 안정적으로 이어지는가?"),
        ("재미·시장성", "밤의 이동·전투·교란이 아침 토론을 풍부하게 하고 반복 플레이 동기가 되는가?"),
        ("완성·출시 가능성", "1인 개발 범위에서 출시 전 우선해야 할 밸런스·UX·운영 과제는 무엇인가?"),
    ], [2700, 6660], font_size=8.9)
    add_heading(doc, "지원 활용", level=2)
    for text in [
        "최대 12인 실제 플레이테스트를 반복할 수 있는 참가자·전시·시연 환경 확보",
        "역할 간 정보량과 직무·안정도 수치에 대한 전문가 멘토링",
        "한 문장과 짧은 영상으로 차별점을 전달하는 피칭·마케팅 메시지 고도화",
        "Steam 출시 전 네트워크 안정성·온보딩·UX의 우선순위 정립",
    ]:
        bullet(doc, text, bullet_id)
    add_callout(doc, "이번 오디션의 목표", "밤의 행동이 다음 날 추리가 되는지 6~12인 플레이에서 검증하고, 출시 가능한 제품으로 다듬습니다.")

    page_break(doc)
    add_page_title(doc, 13, "회사·팀 소개", "[[TEAM_NAME]]은 기획과 프로그래밍을 하나의 플레이 가능한 빌드로 연결해 온 1인 개발팀입니다.")
    add_table(doc, ["항목", "내용"], [
        ("팀명", "[[TEAM_NAME]]"),
        ("개발 형태", "일반부문 개인 참가 · 1인 개발"),
        ("개발자", "[[DEVELOPER_NAME]]"),
        ("주요 역할", "게임 기획 · Unity/C# 프로그래밍 · 네트워크 · UI·게임플레이 통합"),
        ("주요 기술", "Unity · C# · Netcode for GameObjects · Unity Relay · Vivox · Git/GitHub"),
        ("연락처", "[[CONTACT_EMAIL]]"),
    ], [2300, 7060], font_size=9.1)
    add_heading(doc, "개발자 소개", level=2)
    add_para(doc, "Unity와 C# 기반 게임플레이·멀티플레이 기능을 직접 구현해 왔습니다. 2인 협동 퍼즐 프로젝트의 Photon PUN2 로비·동기화와 게임잼 개인작 「SplitTime」의 배포를 경험했고, 현재는 [[GAME_TITLE]]의 기획부터 NGO·Relay·Vivox 기반 로비, 역할, 전투, 채팅과 UI까지 전체 프로토타입을 개발하고 있습니다.")
    add_heading(doc, "1인 개발의 실행 방식", level=2)
    bullet(doc, "플레이 가능한 단위를 먼저 완성하고, 실제 테스트에서 드러난 문제를 우선 수정합니다.", bullet_id)
    bullet(doc, "기존 에셋은 효율적으로 활용하고, 차별화가 필요한 게임 규칙·네트워크·UX 구현에 집중합니다.", bullet_id)
    add_heading(doc, "프로젝트 수행 강점", level=2)
    add_table(doc, ["강점", "근거"], [
        ("구현력", "기획한 20종 역할을 실제 네트워크 권한·행동·승리 판정까지 연결"),
        ("문제 해결", "접속·동기화·고지연·음성·UI 문제를 재현하고 반복 수정"),
        ("범위 관리", "1인 개발 범위에서 핵심 루프를 유지하고 출시 품질 과제를 우선"),
    ], [2300, 7060], font_size=8.9)
    add_callout(doc, "TEAM DIRECTION", "출시 단계의 아트·사운드·QA 협업을 검토하되, 핵심 규칙과 네트워크 구조는 [[TEAM_NAME]]이 관리합니다.", fill=PALE, accent=TEAL)

    p = add_para(doc, "PROJECT STATEMENT", size=9, bold=True, color=TEAL, align=WD_ALIGN_PARAGRAPH.CENTER, after=18)
    p.paragraph_format.page_break_before = True
    add_para(doc, "정체는 사라져도,\n행동은 흔적으로 남습니다.", size=25, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, after=18)
    add_para(doc, "[[GAME_TITLE]]", size=18, bold=True, color=GOLD, align=WD_ALIGN_PARAGRAPH.CENTER, after=24)
    add_para(doc, "밤의 마을을 직접 걷고, 서로의 집과 문을 넘고, 다음 날 그 흔적을 말로 증명하는 3D 소셜 디덕션 게임.", size=12, color=INK, align=WD_ALIGN_PARAGRAPH.CENTER, after=35)
    add_callout(doc, "CURRENT STATUS", "로비 → 역할 배정 → 아침 토론·투표 → 밤 영체 행동 → 사망·승리 → 결과 공개까지 실제 멀티플레이로 연결된 프로토타입", fill=PALE_GOLD, accent=GOLD)
    add_para(doc, "[[TEAM_NAME]]", size=12, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    add_para(doc, "[[DEVELOPER_NAME]]  ·  [[CONTACT_EMAIL]]", size=9.5, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)

    core = doc.core_properties
    core.title = "2026 경기게임오디션 게임·팀 소개서 - [[GAME_TITLE]]"
    core.subject = "2026 경기게임오디션 일반부문 자유양식 게임·팀 소개서"
    core.author = "[[TEAM_NAME]]"
    core.keywords = "경기게임오디션, 게임소개서, 소셜 디덕션, Unity, 멀티플레이"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
